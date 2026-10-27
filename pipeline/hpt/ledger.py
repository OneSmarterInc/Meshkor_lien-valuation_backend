"""The ledger: one row per hospital (CCN) × period, with its state, the last reason it stopped,
how many attempts, and what a person should do next. Written as CSV and Excel after every run:
  data/reports/hpt_ledger.csv, data/reports/hpt_ledger.xlsx

States (in order):  listed, website, url_found, downloaded, validated, registered, loaded, verified
Special:            unchanged (same file as an earlier period: nothing to load), parked (needs a person)

Reason codes and what the pipeline does with them:
"""
from __future__ import annotations

import csv
import re
from datetime import datetime, timezone
from pathlib import Path

from ..common import config

STATES = ["listed", "website", "url_found", "downloaded", "validated", "registered", "loaded", "verified", "unchanged", "parked"]

# code -> (retry this many times automatically, next action for a person)
CODES = {
    "NO_WEBSITE":   (0, "find the hospital website; add it to config/hospital_sites.csv"),
    "NO_HPT_TXT":   (1, "open the site in Chrome, find the price-transparency page; put its cms-hpt.txt or page URL in hospital_sites.csv"),
    "UNMATCHED":    (0, "cms-hpt.txt lists locations but none matched this hospital name; add location_name for this CCN in hospital_sites.csv"),
    "AUTH":         (0, "the file is behind a login; ask the hospital (contact in hpt_catalog.csv) or download by hand"),
    "BLOCKED":      (3, "site refuses our address (403/429); retried through Chrome; try later or by hand"),
    "NOT_FOUND":    (1, "link is dead (404); re-read cms-hpt.txt next run; if it stays dead, check the page by hand"),
    "TIMEOUT":      (3, "download timed out; retried automatically; if it keeps failing, download by hand in Chrome"),
    "CONN":         (3, "could not connect; retried automatically"),
    "HTML":         (0, "server sent a web page instead of a data file (login or block); open the link in Chrome and see"),
    "BAD_FORMAT":   (0, "file is not CSV/JSON/ZIP/XLSX; open it and tell the developer"),
    "INVALID":      (0, "file downloaded but does not look like a CMS price file (header/columns); tell the developer"),
    "PARSE_ERROR":  (0, "loader could not read the file; quarantined; tell the developer (error in the ledger)"),
    "VERIFY_FAIL":  (0, "loaded but failed a verification check; quarantined; tell the developer"),
}
PARK_AT_ONCE = {c for c, (n, _) in CODES.items() if n == 0}

FIELDS = ["ccn", "hospital_name", "system", "tier", "period", "state", "code", "detail", "attempts", "first_seen",
          "last_try", "file", "sha256", "rows", "file_dated", "verify", "next_action"]


class Ledger:
    def __init__(self, data_root: str | None = None):
        self.paths = config.data_paths(data_root)
        self.csv_path = self.paths.root / "reports" / "hpt_ledger.csv"
        self.rows: dict[tuple[str, str], dict] = {}
        if self.csv_path.exists():
            with open(self.csv_path, encoding="utf-8-sig", newline="") as f:
                for r in csv.DictReader(f):
                    if not re.fullmatch(r"\d{6}", r.get("ccn", "")):      # old twin rows like 36S051: dropped
                        continue
                    self.rows[(r["ccn"], r["period"])] = {k: r.get(k, "") for k in FIELDS}

    def get(self, ccn: str, period: str) -> dict:
        key = (str(ccn), str(period))
        if key not in self.rows:
            self.rows[key] = {k: "" for k in FIELDS}
            self.rows[key].update(ccn=str(ccn), period=str(period), state="listed", attempts="0",
                                  first_seen=datetime.now(timezone.utc).isoformat(timespec="seconds"))
        return self.rows[key]

    def set_state(self, ccn: str, period: str, state: str, **fields) -> dict:
        r = self.get(ccn, period)
        r.update(state=state, code="", detail="", next_action="")
        r.update({k: ("" if v is None else str(v)) for k, v in fields.items()})
        r["last_try"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
        return r

    def fail(self, ccn: str, period: str, code: str, detail: str = "", **fields) -> dict:
        """Record a failure. Returns the row; row['state'] == 'parked' when a person must act."""
        r = self.get(ccn, period)
        retries, action = CODES.get(code, (0, "tell the developer"))
        attempts = int(r.get("attempts") or 0) + 1
        r.update(code=code, detail=detail[:300], attempts=str(attempts), next_action=action,
                 last_try=datetime.now(timezone.utc).isoformat(timespec="seconds"))
        r.update({k: ("" if v is None else str(v)) for k, v in fields.items()})
        if code in PARK_AT_ONCE or attempts > retries:
            r["state"] = "parked"
        return r

    def skip(self, ccn: str, period: str, retry_parked: bool = False) -> str | None:
        """Why this hospital is skipped this run, or None to try it."""
        r = self.rows.get((str(ccn), str(period)))
        if not r:
            return None
        if r["state"] in ("loaded", "verified", "unchanged"):
            return r["state"]
        if r["state"] == "parked" and not retry_parked:
            return f"parked ({r['code']})"
        return None

    def unpark(self, ccn: str, period: str) -> None:
        r = self.get(ccn, period)
        if r["state"] == "parked":
            r.update(state="listed", attempts="0", code="", next_action="")

    def save(self) -> tuple[Path, Path | None]:
        self.csv_path.parent.mkdir(parents=True, exist_ok=True)
        rows = sorted(self.rows.values(), key=lambda r: (r["period"], r.get("tier") or "9", r["ccn"]))
        with open(self.csv_path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=FIELDS); w.writeheader(); w.writerows(rows)
        xlsx = self.csv_path.with_suffix(".xlsx")
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
            wb = Workbook(); ws = wb.active; ws.title = "ledger"
            ws.append(FIELDS)
            for c in ws[1]:
                c.font = Font(bold=True)
            fill = {"parked": "FFC7CE", "verified": "C6EFCE", "loaded": "E2EFDA", "unchanged": "DDEBF7"}
            for r in rows:
                ws.append([r.get(k, "") for k in FIELDS])
                if r["state"] in fill:
                    for c in ws[ws.max_row]:
                        c.fill = PatternFill("solid", fgColor=fill[r["state"]])
            for i, k in enumerate(FIELDS, 1):
                ws.column_dimensions[get_column_letter(i)].width = {"detail": 60, "next_action": 60, "hospital_name": 36, "file": 44}.get(k, 14)
            ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
            s2 = wb.create_sheet("summary")
            s2.append(["state", "hospitals"])
            for st in STATES:
                s2.append([st, sum(1 for r in rows if r["state"] == st)])
            s2.append([]); s2.append(["code", "hospitals", "next action"])
            for code, (_, action) in CODES.items():
                n = sum(1 for r in rows if r["code"] == code)
                if n:
                    s2.append([code, n, action])
            wb.save(xlsx)
        except Exception as e:  # noqa: BLE001  (openpyxl missing: CSV is enough)
            print(f"ledger: xlsx not written ({e}); csv is at {self.csv_path}")
            xlsx = None
        return self.csv_path, xlsx

    def summary(self) -> dict:
        rows = list(self.rows.values())
        out = {st: sum(1 for r in rows if r["state"] == st) for st in STATES}
        out["codes"] = {}
        for r in rows:
            if r["code"]:
                out["codes"].setdefault(r["code"], []).append(f"{r['ccn']} {r['hospital_name']}".strip())
        return out


Ledger.__doc__ = (Ledger.__doc__ or "") + "\n" + "\n".join(f"  {c:12s} retries {n}: {a}" for c, (n, a) in CODES.items())
