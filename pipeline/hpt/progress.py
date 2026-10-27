"""The scoreboard: how many Ohio hospitals are done, how many remain, and how much disk is used.

Reads only what already exists (nothing is downloaded):
  config/hospital_sites.csv       website per hospital (what the data member fills in)
  config/hpt_tiers.csv            which tier a hospital system belongs to (1 = first)
  data/ref/hpt_hospitals.parquet  the full Ohio list from pipeline.hpt.hospitals (if built)
  data/ref/hpt_catalog.parquet    price-file URLs found by pipeline.hpt.discover (if built)
  <raw>/MRF/<period>/<ccn>/       files downloaded (by script or by hand)
  config/sources_local.yaml       releases registered for loading
  data/ref/_manifest/ref_hospital_mrf/<ccn>_<period>.json   loaded rows per hospital
  data/reports/verify_hpt.md      last verify result

Writes data/reports/hpt_progress.md and hpt_progress.csv. One row per hospital, stage:
  0 no website | 1 website known | 2 file URL found | 3 downloaded | 4 loaded | 5 verified
Run after every hospital:  python -m pipeline.hpt.progress
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import yaml

from ..common import config

STAGES = {0: "no website", 1: "website known", 2: "file URL found", 3: "downloaded", 4: "loaded", 5: "verified"}


def _csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def folder_bytes(p: Path) -> int:
    return sum(f.stat().st_size for f in p.rglob("*") if f.is_file()) if p.exists() else 0


def collect(data_root: str | None = None) -> tuple[list[dict], dict]:
    paths = config.data_paths(data_root)
    cfg = config.PROJECT_ROOT / "config"
    sites = {r["ccn"]: r for r in _csv(cfg / "hospital_sites.csv") if r.get("ccn")}
    tiers = {r["system"].strip().lower(): r["tier"] for r in _csv(cfg / "hpt_tiers.csv") if r.get("system")}

    hospitals: dict[str, dict] = {}
    hp = paths.ref / "hpt_hospitals.parquet"
    if hp.exists():
        import polars as pl
        for r in pl.read_parquet(hp).iter_rows(named=True):
            hospitals[str(r["ccn"])] = {"ccn": str(r["ccn"]), "hospital_name": r.get("hospital_name") or r.get("name") or "",
                                        "system": r.get("system") or "", "city": r.get("city") or ""}
    for ccn, r in sites.items():                       # sites with a CCN count even before hospitals.py ran
        hospitals.setdefault(ccn, {"ccn": ccn, "hospital_name": r.get("hospital_name", ""), "system": r.get("system", ""), "city": ""})

    catalog: dict[str, str] = {}
    cp = paths.ref / "hpt_catalog.parquet"
    if cp.exists():
        import polars as pl
        for r in pl.read_parquet(cp).iter_rows(named=True):
            if r.get("mrf_url"):
                catalog[str(r["ccn"])] = r["mrf_url"]

    raw_mrf = paths.raw / "MRF"
    downloaded: dict[str, list[str]] = {}
    raw_bytes: dict[str, int] = {}
    if raw_mrf.exists():
        for per in sorted(p for p in raw_mrf.iterdir() if p.is_dir()):
            for c in (p for p in per.iterdir() if p.is_dir()):
                if any(f.suffix.lower() in (".csv", ".json", ".xlsx", ".zip") for f in c.iterdir() if f.is_file()):
                    downloaded.setdefault(c.name, []).append(per.name)
                    raw_bytes[c.name] = raw_bytes.get(c.name, 0) + folder_bytes(c)

    registered: dict[str, list[str]] = {}
    sl = cfg / "sources_local.yaml"
    if sl.exists():
        y = yaml.safe_load(sl.read_text(encoding="utf-8")) or {}
        for rel in (y.get("ref_hospital_mrf") or {}).get("releases") or []:
            registered.setdefault(str(rel.get("ccn")), []).append(str(rel.get("period")))

    loaded: dict[str, dict] = {}
    mdir = paths.manifests / "ref_hospital_mrf"
    if mdir.exists():
        for mf in mdir.glob("*.json"):
            if mf.name.endswith(".quarantine.json"):
                continue
            m = json.loads(mf.read_text(encoding="utf-8"))
            c = str(m.get("ccn")); e = loaded.setdefault(c, {"periods": [], "rows": 0, "last_updated_on": None, "name": m.get("hospital_name")})
            e["periods"].append(str(m.get("period"))); e["rows"] += int(m.get("rows") or 0)
            e["last_updated_on"] = max(filter(None, [e["last_updated_on"], m.get("last_updated_on")]), default=None)
    parquet_bytes = {c.name.split("_")[0]: folder_bytes(c) for c in (paths.ref / "ref_hospital_mrf").iterdir()} if (paths.ref / "ref_hospital_mrf").exists() else {}

    verified: dict[str, str] = {}
    vr = paths.root / "reports" / "verify_hpt.md"
    if vr.exists():
        for m in re.finditer(r"^\|\s*(\d{6})\s*\|\s*(\d{4}-\d{2})\s*\|[^|]*\|\s*(\d+)\s*\|\s*(\d+)\s*\|", vr.read_text(encoding="utf-8"), re.M):
            verified[m.group(1)] = "pass" if m.group(4) == "0" else f"{m.group(4)} fail"
        if not verified:   # fall back: any PASS/FAIL words per CCN heading
            for sec in re.split(r"\n#+ ", vr.read_text(encoding="utf-8")):
                m = re.search(r"\b(\d{6})\b", sec)
                if m:
                    verified[m.group(1)] = "fail" if re.search(r"\bFAIL\b", sec) else "pass"

    rows = []
    for ccn, h in sorted(hospitals.items()):
        site = sites.get(ccn, {})
        stage = 0
        if site.get("website") or site.get("hpt_txt_url"):
            stage = 1
        if ccn in catalog:
            stage = 2
        if ccn in downloaded:
            stage = 3
        if ccn in loaded:
            stage = 4
        if ccn in loaded and verified.get(ccn) == "pass":
            stage = 5
        system = h.get("system") or site.get("system") or ""
        rows.append({"ccn": ccn, "hospital_name": h.get("hospital_name") or (loaded.get(ccn) or {}).get("name") or "",
                     "system": system, "tier": tiers.get(system.strip().lower(), ""),
                     "stage": stage, "stage_name": STAGES[stage],
                     "website": site.get("website", ""), "file_url": catalog.get(ccn, ""),
                     "periods_downloaded": " ".join(downloaded.get(ccn, [])), "periods_registered": " ".join(registered.get(ccn, [])),
                     "periods_loaded": " ".join((loaded.get(ccn) or {}).get("periods", [])),
                     "rows_loaded": (loaded.get(ccn) or {}).get("rows", 0), "file_dated": (loaded.get(ccn) or {}).get("last_updated_on") or "",
                     "verify": verified.get(ccn, ""), "raw_mb": round(raw_bytes.get(ccn, 0) / 1e6, 1), "parquet_mb": round(parquet_bytes.get(ccn, 0) / 1e6, 1),
                     "notes": site.get("notes", "")})
    totals = {"hospitals": len(rows), "by_stage": {s: sum(1 for r in rows if r["stage"] == s) for s in STAGES},
              "loaded": sum(1 for r in rows if r["stage"] >= 4), "remaining": sum(1 for r in rows if r["stage"] < 4),
              "raw_gb": round(sum(raw_bytes.values()) / 1e9, 2), "parquet_gb": round(sum(parquet_bytes.values()) / 1e9, 2),
              "duckdb_gb": round(paths.duckdb.stat().st_size / 1e9, 2) if paths.duckdb.exists() else 0.0,
              "raw_dir": str(paths.raw), "ref_dir": str(paths.ref), "hospital_list_built": hp.exists()}
    return rows, totals


def write(rows, totals, paths) -> tuple[Path, Path]:
    out_md = paths.root / "reports" / "hpt_progress.md"
    out_csv = paths.root / "reports" / "hpt_progress.csv"
    out_md.parent.mkdir(parents=True, exist_ok=True)
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else ["ccn"]); w.writeheader(); w.writerows(rows)
    t = totals
    L = [f"# Hospital price files: progress ({datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC)", "",
         f"**{t['loaded']} of {t['hospitals']} hospitals loaded, {t['remaining']} remaining.**"
         + ("" if t["hospital_list_built"] else "  (run `python -m pipeline.hpt.hospitals` to get the full Ohio list; this counts only hospital_sites.csv)"), "",
         "| Stage | Hospitals |", "|---|---|"] + [f"| {s} {STAGES[s]} | {n} |" for s, n in t["by_stage"].items()] + [
         "", "| Disk | Size |", "|---|---|", f"| raw files `{Path(t['raw_dir']) / 'MRF'}` | {t['raw_gb']} GB |",
         f"| Parquet `{Path(t['ref_dir']) / 'ref_hospital_mrf'}` | {t['parquet_gb']} GB |", f"| DuckDB | {t['duckdb_gb']} GB |", "",
         "## Per hospital", "", "| Tier | CCN | Hospital | System | Stage | Downloaded | Loaded (rows) | File dated | Verify | Raw MB | Website / file URL |", "|---|---|---|---|---|---|---|---|---|---|---|"]
    for r in sorted(rows, key=lambda r: (r["tier"] or "9", -r["stage"], r["ccn"])):
        L.append(f"| {r['tier']} | {r['ccn']} | {r['hospital_name']} | {r['system']} | {r['stage']} {r['stage_name']} | {r['periods_downloaded']} | "
                 f"{r['periods_loaded']} ({r['rows_loaded']:,}) | {r['file_dated']} | {r['verify']} | {r['raw_mb']} | {r['file_url'] or r['website']} |")
    L += ["", "## Next actions", ""]
    todo = [r for r in rows if r["stage"] < 4]
    for s, text in ((0, "find the website, add to config/hospital_sites.csv"), (1, "run discover (or open <website>/cms-hpt.txt in the browser)"),
                    (2, "download the file (download.py, or by hand in the browser then intake.py)"), (3, "register in sources_local.yaml and load_ref")):
        names = [f"{r['ccn']} {r['hospital_name']}" for r in todo if r["stage"] == s]
        if names:
            L.append(f"- Stage {s} → {text}: " + "; ".join(names[:12]) + (f" … +{len(names)-12}" if len(names) > 12 else ""))
    out_md.write_text("\n".join(L) + "\n", encoding="utf-8")
    return out_md, out_csv


def write_status_xlsx(rows, totals, paths) -> Path | None:
    """data/reports/hpt_status.xlsx: the readable workbook (3 small sheets), built from the same facts."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None
    WORDS = {5: "Done (in database, verified)", 4: "Loaded (not yet verified)", 3: "Downloaded (not yet loaded)",
             2: "Link found (not downloaded)", 1: "Website known (not tried)", 0: "No website known"}
    led = {}
    lp = paths.root / "reports" / "hpt_ledger.csv"
    if lp.exists():
        for r in _csv(lp):
            led[r["ccn"]] = r
    wb = Workbook()
    s1 = wb.active; s1.title = "Summary"
    s1.append(["Hospital price files: status", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")])
    s1.append([]); s1.append(["Hospitals in Ohio", totals["hospitals"]])
    s1.append(["Done (verified in database)", totals["by_stage"].get(5, 0)])
    s1.append(["Loaded, verify pending", totals["by_stage"].get(4, 0)])
    s1.append(["Downloaded, load pending", totals["by_stage"].get(3, 0)])
    s1.append(["Remaining (nothing downloaded yet)", totals["by_stage"].get(0, 0) + totals["by_stage"].get(1, 0) + totals["by_stage"].get(2, 0)])
    s1.append([]); s1.append(["Raw files on disk (GB)", totals["raw_gb"]]); s1.append(["Parquet (GB)", totals["parquet_gb"]]); s1.append(["DuckDB (GB)", totals["duckdb_gb"]])
    s1.append([]); s1.append(["Sheet 'Hospitals' = one row per hospital. Sheet 'To do' = only the ones a person must act on."])
    s1.column_dimensions["A"].width = 40; s1.column_dimensions["B"].width = 22
    s1["A1"].font = Font(bold=True, size=13)

    s2 = wb.create_sheet("Hospitals")
    hdr = ["CCN", "Hospital", "System", "Tier", "Status", "File from hospital", "File dated", "Rows loaded", "Checks", "What to do"]
    s2.append(hdr)
    fill = {5: "C6EFCE", 4: "E2EFDA", 3: "FFF2CC", 2: "FFF2CC", 1: "FFFFFF", 0: "F2F2F2"}
    for r in sorted(rows, key=lambda r: (r["tier"] or "9", -r["stage"], r["ccn"])):
        L = led.get(r["ccn"], {})
        todo = L.get("next_action", "") if L.get("state") == "parked" else ""
        s2.append([r["ccn"], r["hospital_name"], r["system"], r["tier"], WORDS[r["stage"]], r["notes"] if False else (L.get("file") or ""),
                   r["file_dated"], r["rows_loaded"] or None, r["verify"], todo])
        for c in s2[s2.max_row]:
            c.fill = PatternFill("solid", fgColor="FFC7CE" if todo else fill[r["stage"]])
    for i, w in enumerate([9, 42, 22, 5, 30, 48, 11, 12, 8, 60], 1):
        s2.column_dimensions[get_column_letter(i)].width = w
    for c in s2[1]:
        c.font = Font(bold=True)
    s2.freeze_panes = "A2"; s2.auto_filter.ref = s2.dimensions

    s3 = wb.create_sheet("To do")
    s3.append(["CCN", "Hospital", "Reason", "Detail", "What to do"])
    for r in sorted(rows, key=lambda r: (r["tier"] or "9", r["ccn"])):
        L = led.get(r["ccn"], {})
        if L.get("state") == "parked":
            s3.append([r["ccn"], r["hospital_name"], L.get("code", ""), L.get("detail", "")[:120], L.get("next_action", "")])
    for i, w in enumerate([9, 42, 14, 60, 60], 1):
        s3.column_dimensions[get_column_letter(i)].width = w
    for c in s3[1]:
        c.font = Font(bold=True)
    s3.freeze_panes = "A2"; s3.auto_filter.ref = s3.dimensions
    out = paths.root / "reports" / "hpt_status.xlsx"
    wb.save(out)
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data-root", default=None)
    a = ap.parse_args(argv)
    rows, totals = collect(a.data_root)
    md, csvp = write(rows, totals, config.data_paths(a.data_root))
    xl = write_status_xlsx(rows, totals, config.data_paths(a.data_root))
    print(f"{totals['loaded']} of {totals['hospitals']} loaded, {totals['remaining']} remaining | raw {totals['raw_gb']} GB, parquet {totals['parquet_gb']} GB")
    print(f"-> {md}\n-> {csvp}" + (f"\n-> {xl}   (open this one)" if xl else ""))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
