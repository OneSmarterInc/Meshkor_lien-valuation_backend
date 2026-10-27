"""Read a hospital bill into header fields and lines.

Accepted files:
  PDF   with a text layer (printed by the billing system) -> read directly, exact
  PDF   without text (a scan)                             -> each page read by local OCR (ocr.py)
  JPG / PNG / TIFF (a photo or scan)                      -> local OCR
  DOCX  (Word)                                            -> paragraphs and tables; pictures inside it by OCR
Each line must show, in this order:
    [line no]  date  revenue code  [HCPCS/CPT]  units  description  charge
OCR is never exact: every line read by OCR carries the OCR certainty, a missing value is filled with a
safe default and lowers the certainty, and intake adds a check flag so a person compares the lines
with the bill.

Nothing here touches Django or the database, so it can be tested alone.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation


class BillNotReadable(Exception):
    """The file has no text we can read (scan, image, password, broken PDF)."""


@dataclass
class BillLine:
    line_no: int
    service_date: date | None
    revenue_code: str
    hcpcs: str
    modifiers: str
    units: Decimal
    description: str
    charge: Decimal
    confidence: float


@dataclass
class Bill:
    facility_name: str = ""
    ccn: str = ""
    npi: str = ""
    type_of_bill: str = ""
    ms_drg: str = ""
    date_from: date | None = None
    date_to: date | None = None
    payer: str = ""
    reason: str = ""
    stated_total: Decimal | None = None
    lines: list[BillLine] = field(default_factory=list)
    diagnoses: list[str] = field(default_factory=list)
    pages: int = 0
    notes: list[str] = field(default_factory=list)
    method: str = "text"          # text | ocr | docx | docx+ocr
    ocr_score: float | None = None   # lowest OCR certainty of a bill line (OCR only)

    @property
    def lines_total(self) -> Decimal:
        return sum((ln.charge for ln in self.lines), Decimal("0"))

    @property
    def inpatient(self) -> bool:
        tob = self.type_of_bill.lstrip("0")
        return tob[:2] in ("11", "12") or bool(self.ms_drg)


# ----------------------------------------------------------------------------- text

def pdf_text(data: bytes) -> tuple[list[str], int]:
    """Text lines of every page. Large characters (watermarks, logos) are dropped."""
    import pdfplumber
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001 — pdfminer raises many kinds
        raise BillNotReadable(f"The file is not a readable PDF ({type(e).__name__}).") from e
    lines: list[str] = []
    with pdf:
        n = len(pdf.pages)
        for page in pdf.pages:
            page = page.filter(lambda o: o.get("object_type") != "char" or (o.get("size") or 0) < 20)
            text = page.extract_text() or ""
            lines += [ln.strip() for ln in text.splitlines() if ln.strip()]
    return lines, n


def docx_text(data: bytes) -> tuple[list[str], list[bytes]]:
    """Text lines of a Word file (paragraphs, then table rows with cells joined), and its pictures."""
    import zipfile
    try:
        import docx
    except ImportError as e:
        raise BillNotReadable("Word files cannot be read on this server yet: run  python -m pip install -r requirements.txt") from e
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise BillNotReadable(f"The Word file could not be opened ({type(e).__name__}).") from e
    lines = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells, prev = [], None
            for c in row.cells:            # merged cells repeat: keep each once
                if c._tc is not prev:
                    txt = " ".join(c.text.split())
                    if txt:
                        cells.append(txt)
                prev = c._tc
            if cells:
                lines.append(" ".join(cells))
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        pics = [z.read(n) for n in z.namelist() if n.startswith("word/media/") and n.lower().endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff"))]
    return lines, pics


# ----------------------------------------------------------------------------- parsing

DATE = r"(\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{2,4})"
MONEY = r"\(?-?\$?\s?-?[\d,]+\.\d{2}\)?"
LINE_RX = re.compile(
    rf"^(?:(?P<no>\d{{1,4}})\s+)?(?P<date>{DATE})\s+(?P<rev>\d{{3,4}})\s+"
    rf"(?:(?P<code>[A-Z]\d{{4}}|\d{{4}}[A-Z0-9])(?:[\s-]+(?P<mods>(?:[A-Z0-9]{{2}}\s?){{1,4}}))?\s+)?"
    rf"(?P<units>\d{{1,4}}(?:\.\d{{1,2}})?)\s+(?P<desc>.+?)\s+(?P<charge>{MONEY})$")
# A looser pattern, tried only when LINE_RX does not match: OCR often drops a thin "1" in the units column,
# or a revenue code or the description. The missing part is filled with a safe default and the line gets a
# lower certainty, so the reviewer is asked to check it.
LOOSE_RX = re.compile(
    rf"^(?:(?P<no>\d{{1,4}})\s+)?(?P<date>{DATE})\s+(?:(?P<rev>\d{{3,4}})\s+)?"
    rf"(?:(?P<code>[A-Z]\d{{4}}|\d{{4}}[A-Z0-9])(?:[\s-]+(?P<mods>(?:[A-Z0-9]{{2}}\s?){{1,4}}))?\s+)?"
    rf"(?:(?P<units>\d{{1,4}}(?:\.\d{{1,2}})?)\s+)?(?P<desc>.*?)\s*(?P<charge>{MONEY})$")
TOTAL_RX = re.compile(rf"^[\W_]*(?:TOTAL CHARGES|TOTAL BILLED|TOTAL|GRAND TOTAL|TOTALS)\b[:\s]*(?P<amt>{MONEY})$", re.I)
HEADER_RX = re.compile(r"^#?\s*(?:line\s+)?date\b.*\b(rev|revenue)\b", re.I)

LABELS = {
    "facility_name": r"(?:Facility|Hospital|Provider)(?: name)?",
    "ccn": r"(?:CCN|Medicare (?:provider )?(?:number|no\.?)|Provider number)",
    "npi": r"(?:Facility NPI|Billing NPI|NPI)",
    "type_of_bill": r"(?:Type of bill|TOB|Bill type)",
    "ms_drg": r"(?:MS-DRG|DRG)",
    "payer": r"(?:Primary payer|Payer|Insurance|Primary insurance|Plan)",
    "reason": r"(?:Reason for visit|Reason|Visit reason)",
    "dates": r"(?:Service dates|Statement covers period|Dates of service|Service period|From)",
}
STOP = r"(?=\s+(?:Statement no\.|Account no\.|(?:Facility NPI|NPI|MS-DRG|DRG|Primary payer|Payer|Patient)\b)|$)"


def _date(s: str | None) -> date | None:
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _money(s: str) -> Decimal:
    neg = s.strip().startswith("(") or "-" in s
    v = Decimal(re.sub(r"[^\d.]", "", s))
    return -v if neg else v


def _field(text: str, label: str) -> str:
    m = re.search(rf"(?:^|\s){label}[:#]?\s+(.+?){STOP}", text, re.I)
    return m.group(1).strip() if m else ""


def parse_header(lines: list[str], bill: Bill) -> None:
    head = []
    for ln in lines:
        if HEADER_RX.match(ln) or LINE_RX.match(ln):
            break
        head.append(ln)
    for ln in head:
        if not bill.facility_name:
            bill.facility_name = _field(ln, LABELS["facility_name"])
        if not bill.ccn:
            m = re.search(rf"{LABELS['ccn']}[:#]?\s+(\d{{2}}[0-9A-Z]\d{{3}})\b", ln, re.I)
            bill.ccn = m.group(1) if m else ""
        if not bill.npi:
            m = re.search(rf"{LABELS['npi']}[:#]?\s+(\d{{10}})\b", ln, re.I)
            bill.npi = m.group(1) if m else ""
        if not bill.type_of_bill:
            m = re.search(rf"{LABELS['type_of_bill']}[:#]?\s+(0?\d{{3}})\b", ln, re.I)
            bill.type_of_bill = m.group(1).zfill(4) if m else ""
        if not bill.ms_drg:
            m = re.search(rf"{LABELS['ms_drg']}[:#]?\s+(\d{{3}})\b", ln, re.I)
            bill.ms_drg = m.group(1) if m else ""
        if not bill.payer:
            bill.payer = _field(ln, LABELS["payer"])
        if not bill.reason:
            bill.reason = _field(ln, LABELS["reason"])
        if not bill.date_from:
            m = re.search(rf"{LABELS['dates']}[:#]?\s+{DATE}(?:\s*(?:to|-|through|thru)\s*{DATE})?", ln, re.I)
            if m:
                bill.date_from, bill.date_to = _date(m.group(1)), _date(m.group(2) or m.group(1))
    if bill.facility_name.startswith("("):
        bill.facility_name = ""


DX_LABEL = re.compile(r"\b(?P<kind>principal|primary|other|secondary|additional)?\s*(?:diagnos[ie]s|dx|icd-?10(?:-cm)?(?: codes?)?)\b(?P<rest>.*)$", re.I)
ICD10 = re.compile(r"\b([A-TV-Z]\d[0-9A-Z])(?:\.?([0-9A-Z]{1,4}))?\b")


def parse_diagnoses(lines: list[str], bill: Bill) -> None:
    """ICD-10-CM codes from lines labelled Diagnosis / Dx / ICD-10 (anywhere on the bill: a UB-04 lists
    them below the lines, box 67). The principal diagnosis comes first; "admitting" diagnoses are skipped."""
    principal, others = [], []
    for ln in lines:
        if re.search(r"\badmitting\b", ln, re.I) or LINE_RX.match(ln):
            continue
        m = DX_LABEL.search(ln)
        if not m:
            continue
        codes = [(a + (b or "")).upper() for a, b in ICD10.findall(m.group("rest"))]
        (principal if (m.group("kind") or "").lower() in ("principal", "primary") else others).extend(codes)
    seen = []
    for c in principal + others:
        if c not in seen:
            seen.append(c)
    bill.diagnoses = seen


def parse_lines(lines: list[str], bill: Bill) -> None:
    n = 0
    in_table = False
    for ln in lines:
        if HEADER_RX.match(ln):
            in_table = True
            continue
        m = LINE_RX.match(ln)
        loose = False
        if not m:
            m = LOOSE_RX.match(ln)
            loose = m is not None
        if m:
            in_table = True
            n += 1
            code = (m.group("code") or "").upper()
            rev = (m.group("rev") or "").zfill(4) if m.group("rev") else ""
            conf = 0.98
            missing = []
            if loose:
                if not m.group("units"):
                    missing.append("units (1 assumed)")
                if not m.group("rev"):
                    missing.append("revenue code")
                if not m.group("desc").strip():
                    missing.append("description")
                conf = 0.85 if missing else 0.9
            try:
                units = Decimal(m.group("units")) if m.group("units") else Decimal("1")
                charge = _money(m.group("charge"))
            except InvalidOperation:
                units, charge, conf = Decimal("1"), Decimal("0"), 0.4
            if units <= 0 or units > 999:
                conf = min(conf, 0.7)
            if m.group("no") and int(m.group("no")) != n:
                conf = min(conf, 0.8)   # a line number was skipped: a line may be missing
            bill.lines.append(BillLine(
                line_no=n, service_date=_date(m.group("date")), revenue_code=rev, hcpcs=code,
                modifiers=" ".join((m.group("mods") or "").split()), units=units,
                description=m.group("desc").strip()[:300], charge=charge, confidence=conf))
            if missing:
                bill.notes.append(f"Line {n}: not read: {', '.join(missing)}")
            continue
        t = TOTAL_RX.match(ln)
        if t and in_table:
            bill.stated_total = _money(t.group("amt"))
            in_table = False
            continue
        # A wrapped description: short upper-case text right after a line
        if in_table and bill.lines and len(ln) <= 40 and not re.search(r"\d{2,}", ln) and ln.upper() == ln:
            last = bill.lines[-1]
            last.description = f"{last.description} {ln}"[:300]


def _finish(bill: Bill, lines: list[str], scores: list[float] | None = None) -> Bill:
    parse_header(lines, bill)
    parse_lines(lines, bill)
    parse_diagnoses(lines, bill)
    if not bill.lines:
        raise BillNotReadable("No bill lines were found. The layout is not one the reader knows yet."
                              + (" The picture may be too blurred or too small; upload a sharper photo or the PDF from the hospital."
                                 if scores is not None else ""))
    if scores is not None:
        # OCR: a bill line's certainty is at most the OCR certainty of its text line, and never above 0.95.
        # parse_lines adds one bill line per matching text line, in order, so the two lists pair up.
        matched = [sc for t, sc in zip(lines, scores) if LINE_RX.match(t) or LOOSE_RX.match(t)]
        for ln, sc in zip(bill.lines, matched):
            ln.confidence = round(min(ln.confidence, sc, 0.95), 2)
        bill.ocr_score = min(ln.confidence for ln in bill.lines)
    dates = [ln.service_date for ln in bill.lines if ln.service_date]
    if dates:
        bill.date_from = bill.date_from or min(dates)
        bill.date_to = bill.date_to or max(dates)
    return bill


def _ocr(ocr_lines, bill: Bill) -> Bill:
    bill.method = "ocr" if bill.method == "text" else bill.method
    return _finish(bill, [o.text for o in ocr_lines], [o.score for o in ocr_lines])


def read_pdf(data: bytes) -> Bill:
    lines, pages = pdf_text(data)
    bill = Bill(pages=pages)
    if sum(len(x) for x in lines) >= 40:
        try:
            return _finish(bill, lines)
        except BillNotReadable:
            if sum(len(x) for x in lines) > 400:
                raise              # a real text layer with an unknown layout: OCR would not help
    from . import ocr               # no (usable) text layer: a scanned PDF
    try:
        ol, bill.pages = ocr.pdf_lines(data)
    except ocr.OcrNotInstalled:
        raise
    except Exception as e:  # noqa: BLE001 — pdfium raises its own error types
        raise BillNotReadable(f"The PDF has no text and its pages could not be rendered for OCR ({type(e).__name__}).") from e
    return _ocr(ol, bill)


def read_image(data: bytes) -> Bill:
    from . import ocr
    try:
        ol = ocr.image_bytes_lines(data)
    except ocr.OcrNotInstalled:
        raise
    except Exception as e:  # noqa: BLE001 — PIL raises many kinds
        raise BillNotReadable(f"The picture could not be opened ({type(e).__name__}).") from e
    return _ocr(ol, Bill(pages=1))


def read_docx(data: bytes) -> Bill:
    lines, pics = docx_text(data)
    bill = Bill(pages=1, method="docx")
    try:
        return _finish(bill, lines)
    except BillNotReadable:
        if not pics:
            raise
    from . import ocr               # a Word file that holds a scanned bill as a picture
    ol = []
    for p in pics:
        ol += ocr.image_bytes_lines(p)
    bill = Bill(pages=len(pics), method="docx+ocr")
    return _finish(bill, lines + [o.text for o in ol], [1.0] * len(lines) + [o.score for o in ol])


KINDS = {"pdf": "PDF", "png": "image", "jpg": "image", "tiff": "image", "docx": "Word"}


def kind_of(data: bytes, name: str = "") -> str | None:
    """The real type of a file, from its first bytes (the name can be wrong)."""
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "png"
    if data[:3] == b"\xff\xd8\xff":
        return "jpg"
    if data[:4] in (b"II*\x00", b"MM\x00*"):
        return "tiff"
    if data[:4] == b"PK\x03\x04":
        import zipfile
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            return None
    return None


def read(data: bytes, name: str) -> Bill:
    k = kind_of(data, name)
    if k == "pdf":
        return read_pdf(data)
    if k in ("png", "jpg", "tiff"):
        return read_image(data)
    if k == "docx":
        return read_docx(data)
    raise BillNotReadable("The file is not a PDF, JPG, PNG, TIFF or Word (.docx) file.")
