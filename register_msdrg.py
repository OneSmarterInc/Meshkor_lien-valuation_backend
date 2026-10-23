#!/usr/bin/env python
"""
register_msdrg.py: add the loaded MS-DRG versions to data_register.xlsx (append only, no overwrite).

Usage (from D:\\meshkor-data):
    python register_msdrg.py --who Hansavahini            # every version with a manifest
    python register_msdrg.py --who Hansavahini --version V44
"""
import argparse, json, sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

LOG_HDR = ["Date", "Who", "File / item", "Source page or link", "Raw folder", "Note"]
HDR = ["Version", "Fiscal year", "Effective from", "Effective to", "Table", "Rows", "Raw file", "SHA-256", "Parquet", "Loaded at", "Verify"]


def sheet(wb, name, header):
    if name in wb.sheetnames:
        ws = wb[name]
        if ws.max_row == 1 and ws.cell(1, 1).value is None:
            ws.append(header)
        return ws
    ws = wb.create_sheet(name); ws.append(header); return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--who", required=True)
    ap.add_argument("--version", default="")
    ap.add_argument("--data", default="data")
    ap.add_argument("--xlsx", default="data_register.xlsx")
    a = ap.parse_args()
    man_dir = Path(a.data) / "ref" / "_manifest" / "ref_msdrg"
    mans = [man_dir / f"{a.version}.json"] if a.version else sorted(man_dir.glob("V*.json"))
    if not mans or not mans[0].exists():
        sys.exit(f"no manifests in {man_dir}")
    rep = Path(a.data) / "reports" / "verify_msdrg.md"
    verify = "not run"
    if rep.exists():
        line = [l for l in rep.read_text(encoding="utf-8").splitlines() if "Result:" in l]
        verify = line[0].split("Result:")[1].strip(" *") if line else "see report"

    x = Path(a.xlsx)
    wb = load_workbook(x) if x.exists() else Workbook()
    if not x.exists():
        wb.remove(wb.active)
    log = sheet(wb, "Download log", LOG_HDR)
    ws = sheet(wb, "MS-DRG", HDR)
    have = {str(r[7]) for r in ws.iter_rows(min_row=2, values_only=True) if r and r[7]}
    added = 0
    for mp in mans:
        m = json.loads(mp.read_text()); v = m["version"]
        a_sha = next(s["sha256"] for s in m["sources"] if s["file"] == "appendix_A.txt")
        if a_sha in have:
            print(f"{v}: already registered"); continue
        for s in m["sources"]:
            log.append([date.today().isoformat(), a.who, f"{v}/{s['file']}", m["page_url"], m["raw_folder"],
                        f"MS-DRG {v} " + (s.get("role") or "definitions manual text, loaded")])
        for t, i in m["tables"].items():
            ws.append([v, m["fiscal_year"], m["effective_start"], m["effective_end"], t, i["rows"], i["source_file"],
                       next(s["sha256"] for s in m["sources"] if s["file"] == i["source_file"]),
                       i["output"], m["loaded_at"], verify])
        added += 1
    wb.save(x)
    print(f"registered {added} version(s) in '{x}' (sheets: Download log, MS-DRG)")


if __name__ == "__main__":
    main()
