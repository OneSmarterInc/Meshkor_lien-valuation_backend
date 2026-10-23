#!/usr/bin/env python
"""
register_mrf.py: add one loaded hospital price file to data_register.xlsx.

Appends a row to the "Download log" sheet and a row to a "Hospital MRF" sheet (created if missing).
Never overwrites: a file whose SHA-256 is already in the sheet is skipped.

Usage (from D:\\meshkor-data):
    python register_mrf.py --ccn 360085 --period 2026-10 --who Hansavahini
    python register_mrf.py --ccn 360085 --period 2026-10 --who Hansavahini --xlsx "E:\\Hansavahini\\data_register.xlsx"
"""
import argparse, json, sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

TABLE = "ref_hospital_mrf"
LOG_HDR = ["Date", "Who", "File / item", "Source page or link", "Raw folder", "Note"]
MRF_HDR = ["CCN", "Hospital", "Period", "File dated", "Template", "Rows", "Payers", "Raw file",
           "SHA-256", "Parquet", "Loaded at", "Verify", "Warnings"]


def sheet(wb, name, header):
    if name in wb.sheetnames:
        ws = wb[name]
        if ws.max_row == 1 and ws.cell(1, 1).value is None:
            ws.append(header)
        return ws
    ws = wb.create_sheet(name)
    ws.append(header)
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ccn", required=True)
    ap.add_argument("--period", required=True)
    ap.add_argument("--who", required=True)
    ap.add_argument("--data", default="data")
    ap.add_argument("--xlsx", default="data_register.xlsx")
    a = ap.parse_args()

    man = Path(a.data) / "ref" / "_manifest" / TABLE / f"{a.ccn}_{a.period}.json"
    if not man.exists():
        sys.exit(f"manifest not found: {man}  (run load_mrf.py first)")
    m = json.loads(man.read_text())
    rep = Path(a.data) / "reports" / f"verify_mrf_{a.ccn}_{a.period}.md"
    verify = "not run"
    if rep.exists():
        line = [l for l in rep.read_text(encoding="utf-8").splitlines() if "Result:" in l]
        verify = line[0].split("Result:")[1].strip(" *") if line else "see report"

    x = Path(a.xlsx)
    wb = load_workbook(x) if x.exists() else Workbook()
    if not x.exists():
        wb.remove(wb.active)
    log = sheet(wb, "Download log", LOG_HDR)
    mrf = sheet(wb, "Hospital MRF", MRF_HDR)

    src = m["sources"][0]
    have = {str(r[8]) for r in mrf.iter_rows(min_row=2, values_only=True) if r and r[8]}
    if src["sha256"] in have:
        print(f"already registered (same SHA-256): {src['file']}")
        return
    log.append([date.today().isoformat(), a.who, src["file"], m.get("page_url", ""), m["raw_folder"],
                f"{m['hospital_name']} price file, CCN {a.ccn}, loaded to {TABLE}"])
    mrf.append([a.ccn, m["hospital_name"], a.period, m["last_updated_on"], m["mrf_version"], m["rows"],
                m["distinct_payers"], src["file"], src["sha256"], m["output"], m["loaded_at"], verify,
                "; ".join(m.get("warnings", []))])
    wb.save(x)
    print(f"added {src['file']} to '{x}' (sheets: Download log, Hospital MRF)")


if __name__ == "__main__":
    main()
