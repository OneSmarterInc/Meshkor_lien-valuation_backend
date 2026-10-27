r"""Pricing step tests (pipeline/lien). Run:
    set DB_ENGINE=sqlite && python backend\manage.py test cases accounts
Uses the synthetic test bills in tests/fixtures/test_bills and a tiny DuckDB (lien_testdb.py)."""
import os
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest import mock

from cryptography.fernet import Fernet
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, override_settings
from rest_framework.test import APITestCase

from accounts.models import Firm, Profile
from pipeline.lien import db, match, price, reader
from . import lien_testdb as T
from .models import Case

BILLS = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "test_bills"
TMP = tempfile.mkdtemp()
DATA = Path(tempfile.mkdtemp())
T.build(DATA)


def bill(name):
    return (BILLS / f"{name}_TEST_bill.pdf").read_bytes()


EOBS = BILLS.parent / "test_eobs"


def eob(name):
    return (EOBS / f"{name}_TEST_eob.pdf").read_bytes()


@mock.patch.dict(os.environ, {"MESHKOR_DATA_ROOT": str(DATA)})
class ReaderAndPriceTests(SimpleTestCase):
    def test_reader_gets_header_lines_and_total(self):
        b = reader.read(bill("B01"), "b.pdf")
        self.assertEqual((b.ccn, b.type_of_bill, b.payer), ("360085", "0131", "Anthem"))
        self.assertEqual(b.facility_name, "OHIO STATE UNIVERSITY HOSPITALS")
        self.assertEqual((b.date_from, b.date_to), (date(2026, 8, 14), date(2026, 8, 14)))
        self.assertEqual(len(b.lines), 11)
        self.assertEqual(b.lines_total, Decimal("9647.87")); self.assertEqual(b.stated_total, Decimal("9647.87"))
        ln = b.lines[9]
        self.assertEqual((ln.revenue_code, ln.hcpcs, ln.units, ln.charge), ("0636", "J2405", Decimal("4"), Decimal("1.00")))
        self.assertFalse(b.inpatient)
        b8 = reader.read(bill("B08"), "b.pdf")
        self.assertTrue(b8.inpatient); self.assertEqual(b8.ms_drg, "481")
        self.assertEqual(b8.lines[3].hcpcs, ""); self.assertEqual(b8.lines[3].units, Decimal("3"))   # room and board, no code
        self.assertTrue(b8.lines[11].description.endswith("MINS"))                                   # wrapped text joined

    def _photo(self, name="B01"):
        """A tilted, blurred phone-photo of a synthetic test bill, as JPEG bytes."""
        import io as _io
        import pypdfium2 as pdfium
        from PIL import ImageFilter
        img = pdfium.PdfDocument(bill(name))[0].render(scale=200 / 72).to_pil().convert("RGB")
        img = img.resize((int(img.width * 0.75), int(img.height * 0.75))).rotate(1.0, expand=True, fillcolor=(235, 232, 225))
        img = img.filter(ImageFilter.GaussianBlur(0.6))
        out = _io.BytesIO(); img.save(out, "JPEG", quality=70)
        return out.getvalue()

    def test_ocr_reads_a_phone_photo(self):
        b = reader.read(self._photo("B01"), "photo.jpg")
        self.assertEqual(b.method, "ocr")
        self.assertEqual((b.ccn, b.payer, str(b.date_from)), ("360085", "Anthem", "2026-08-14"))
        self.assertEqual(len(b.lines), 11)
        self.assertEqual(b.lines_total, Decimal("9647.87")); self.assertEqual(b.stated_total, Decimal("9647.87"))
        self.assertTrue(all(ln.confidence <= 0.95 for ln in b.lines))          # OCR is never "certain"
        codes = [ln.hcpcs for ln in b.lines]
        self.assertEqual(codes[:3], ["99285", "70450", "72125"])

    def test_ocr_reads_a_scanned_pdf(self):
        import io as _io
        from PIL import Image
        buf = _io.BytesIO(); Image.open(_io.BytesIO(self._photo("B08"))).save(buf, "PDF")
        b = reader.read(buf.getvalue(), "scan.pdf")
        self.assertEqual(b.method, "ocr"); self.assertTrue(b.inpatient); self.assertEqual(b.ms_drg, "481")
        self.assertEqual(b.lines_total, Decimal("58612.00"))

    def test_missing_ocr_packages_say_so(self):
        import sys
        from pipeline.lien import ocr
        with mock.patch.dict(sys.modules, {"numpy": None}):
            with self.assertRaises(ocr.OcrNotInstalled) as cm:
                reader.read(self._photo("B01"), "photo.jpg")
        self.assertIn("pip install -r requirements.txt", str(cm.exception))

    def test_word_file_with_a_table(self):
        import io as _io
        import docx
        d = docx.Document()
        for t in ("Facility OHIO STATE UNIVERSITY HOSPITALS", "CCN 360085", "Type of bill 0131",
                  "Service dates 2026-09-18 to 2026-09-18", "Primary payer Aetna"):
            d.add_paragraph(t)
        rows = [("#", "Date", "Rev", "HCPCS/CPT", "Units", "Description", "Charge"),
                ("1", "2026-09-18", "0610", "73721", "1", "HC MRI LOWER EXTREMITY JOINT W/O CONTRAST", "$2,957.00")]
        tb = d.add_table(rows=0, cols=7)
        for r in rows:
            cells = tb.add_row().cells
            for c, v in zip(cells, r):
                c.text = v
        d.add_paragraph("TOTAL CHARGES $2,957.00")
        buf = _io.BytesIO(); d.save(buf)
        b = reader.read(buf.getvalue(), "bill.docx")
        self.assertEqual((b.method, b.ccn, b.payer, len(b.lines)), ("docx", "360085", "Aetna", 1))
        ln = b.lines[0]
        self.assertEqual((ln.revenue_code, ln.hcpcs, ln.units, ln.charge), ("0610", "73721", Decimal("1"), Decimal("2957.00")))

    def test_not_a_readable_bill(self):
        with self.assertRaises(reader.BillNotReadable):
            reader.read(b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF", "b.pdf")
        with self.assertRaises(reader.BillNotReadable):
            reader.read(b"\x89PNG....", "scan.png")

    def test_outpatient_line_values(self):
        b = reader.read(bill("B01"), "b.pdf")
        con = db.connect()
        r = price.price(con, b, "360085", b.payer)
        con.close()
        self.assertEqual(r.payer_group, "Anthem"); self.assertTrue(r.price_file.covers_service_date)
        er = r.lines["service_date"][0]
        self.assertEqual(er["payer_rate"], Decimal("1450.00"))          # median of the 2 Anthem outpatient plans; inpatient row ignored
        self.assertEqual(er["median"], Decimal("1450.00"))              # commercial only: 1300, 1400, 1500, 1700 (Humana MA 620, Medicaid 450 left out)
        self.assertEqual(er["medicare"], Decimal(str(T.opps_adjusted(600))))
        self.assertEqual(er["cost"], Decimal("714.63"))                 # 3403 x 0.21
        ct = r.lines["service_date"][1]
        self.assertTrue(ct["payer_rate_is_estimate"]); self.assertEqual(ct["payer_rate"], Decimal("825.60"))   # 40 % of 2064
        self.assertEqual(ct["median"], Decimal("750.00"))                 # Cigna's 0.62 "dollar" row is a factor: left out
        self.assertFalse(price.is_multiplier({"negotiated_dollar": 3.0, "gross_charge": 60.0}, 3.0))   # a real cheap price (venipuncture)
        self.assertTrue(price.is_multiplier({"negotiated_dollar": 0.62, "gross_charge": 3403.0}, 586.78))
        self.assertIn("Factors posted as dollars", " ".join(f["title"] for f in r.flags))
        lab = r.lines["service_date"][5]
        self.assertEqual((lab["medicare"], lab["medicare_source"]), (Decimal("7.50"), "CLFS 2026Q3"))
        self.assertEqual(lab["cash"], Decimal("15.00"))
        xr = r.lines["service_date"][3]                                  # 226.10 % of Medicare, and a $200 plan
        self.assertEqual(xr["payer_rate"], Decimal(str(round((2.261 * T.opps_adjusted(70) + 200) / 2, 2))))
        self.assertTrue(xr["payer_rate_is_estimate"])
        self.assertIn("226.1 %", " ".join(f["title"] for f in r.flags))
        drug = r.lines["service_date"][8]
        self.assertEqual((drug["medicare"], drug["medicare_status"]), (Decimal("0.00"), "packaged"))
        self.assertNotIn("estimate", r.lines)
        self.assertEqual(r.totals["billed"], 9647.87)
        self.assertEqual(r.totals["coverage"]["payer_rate"], round(100 * (3403 + 2064 + 323 + 131) / 9647.87))
        kinds = {f["kind"] for f in r.flags}
        self.assertIn("packaged", kinds); self.assertIn("no_rate", kinds)

    def test_comprehensive_apc_packages_the_other_lines(self):
        b = reader.read(bill("B07"), "b.pdf")
        con = db.connect()
        r = price.price(con, b, "360085", b.payer)
        con.close()
        lines = r.lines["service_date"]
        # anesthesia, recovery room, implant: no code, paid inside the J1 payment -> $0 and counted as covered
        self.assertEqual([(l["medicare"], l["medicare_status"]) for l in lines[1:4]], [(Decimal("0.00"), "packaged")] * 3)
        self.assertEqual(r.totals["coverage"]["medicare"], 100)
        self.assertEqual(r.totals["coverage"]["medicare_billed"], 27671.34)
        self.assertEqual(lines[0]["medicare"], Decimal(str(T.opps_adjusted(6800))))
        self.assertEqual(lines[4]["medicare_status"], "packaged")           # splint 29125 rides on the J1 line
        self.assertEqual(r.totals["medicare"], T.opps_adjusted(6800))

    def _with(self, name, extra):
        """A synthetic test bill plus extra lines: (date, rev, code, units, charge)."""
        b = reader.read(bill(name), "b.pdf")
        n = len(b.lines)
        for i, (dt, rev, code, units, charge) in enumerate(extra, start=1):
            b.lines.append(reader.BillLine(n + i, dt, rev, code, "", Decimal(units), f"TEST {code}", Decimal(charge), 0.98))
        b.date_to = max(ln.service_date for ln in b.lines)
        return b

    def _price(self, b):
        con = db.connect()
        try:
            return price.price(con, b, "360085", b.payer)
        finally:
            con.close()

    def test_capc_bundles_only_its_own_visit_and_keeps_the_exclusions(self):
        later = date(2026, 9, 2)                                   # B07 surgery is 2026-08-20
        b = self._with("B07", [(date(2026, 8, 20), "0278", "C1749", 1, 900),    # pass-through device, same visit
                               (date(2026, 8, 20), "0306", "87070", 1, 80),     # lab on the lab fee schedule (status A)
                               (later, "0420", "97110", 2, 300)])               # therapy two weeks later: its own visit
        r = self._price(b)
        L = {l["hcpcs"]: l for l in r.lines["service_date"] if l["hcpcs"]}
        self.assertEqual(L["25607"]["medicare"], Decimal(str(T.opps_adjusted(6800))))     # primary J1
        self.assertEqual(L["29125"]["medicare_status"], "packaged")                        # splint: inside the J1
        self.assertEqual(L["C1749"]["medicare"], Decimal(str(T.opps_adjusted(400))))       # status H: never packaged
        self.assertEqual(L["87070"]["medicare"], Decimal("11.00"))                         # lab fee schedule: separate
        self.assertEqual(L["97110"]["medicare"], Decimal("56.00"))                         # 2 units x $28, own visit
        self.assertIn("2 visits", " ".join(f["title"] for f in r.flags))

    def test_second_j1_in_the_same_visit_is_packaged(self):
        b = self._with("B07", [(date(2026, 8, 20), "0360", "29881", 1, 9000)])
        r = self._price(b)
        L = {l["hcpcs"]: l for l in r.lines["service_date"] if l["hcpcs"]}
        self.assertIn("primary comprehensive", L["25607"]["medicare_source"])            # 6800 beats 5000
        self.assertEqual((L["29881"]["medicare"], L["29881"]["medicare_status"]), (Decimal("0.00"), "packaged"))
        self.assertEqual(r.totals["medicare"], T.opps_adjusted(6800))

    def test_multiple_procedure_discount_on_t_codes(self):
        b = self._with("B01", [(date(2026, 8, 14), "0450", "12002", 1, 900), (date(2026, 8, 14), "0450", "11042", 2, 700)])
        r = self._price(b)
        L = {l["hcpcs"]: l for l in r.lines["service_date"] if l["hcpcs"]}
        self.assertEqual(L["12002"]["medicare"], Decimal(str(T.opps_adjusted(500))))                    # highest: full
        self.assertEqual(L["11042"]["medicare"], Decimal(str(round(T.opps_adjusted(300) * 2 * 0.5, 2))))  # both units at 50 %
        self.assertIn("multiple procedure discount", L["11042"]["medicare_source"])
        self.assertEqual(L["99285"]["medicare"], Decimal(str(T.opps_adjusted(600))))                    # status V: no discount

    def test_diagnoses_are_read_and_named(self):
        b = reader.Bill()
        reader.parse_diagnoses(["Admitting diagnosis: R51.9",
                                "Other diagnoses: E11.65, D62, J96.01",
                                "Principal diagnosis: S72.301A",
                                "1 2026-07-12 0450 99285 1 HC EMERGENCY ROOM LEVEL 5 $3,403.00"], b)
        self.assertEqual(b.diagnoses, ["S72301A", "E1165", "D62", "J9601"])      # principal first, admitting skipped
        b8 = reader.read(bill("B08"), "b.pdf")
        b8.diagnoses = b.diagnoses
        r = self._price(b8)
        dx = {d["code"]: d for d in r.stay["diagnoses"]}
        self.assertEqual(dx["S72301A"]["level"], "principal")
        self.assertTrue(dx["S72301A"]["description"].startswith("Unspecified fracture of shaft of right femur"))
        self.assertEqual((dx["E1165"]["level"], dx["D62"]["level"], dx["J9601"]["level"]), ("cc", "cc", "mcc"))
        drg = next(f for f in r.flags if f["kind"] == "drg_check")
        self.assertIn("MCC: J9601", drg["detail"]); self.assertIn("CC: E1165, D62", drg["detail"])

    def test_inpatient_stay_priced_as_one_drg(self):
        b = reader.read(bill("B08"), "b.pdf")
        con = db.connect()
        r = price.price(con, b, "360085", b.payer)
        con.close()
        v = r.stay["values"]
        self.assertEqual(r.stay["ms_drg"], "481"); self.assertEqual(r.stay["los_days"], 3)
        self.assertEqual(v["payer_rate"], 22000.0)                          # UnitedHealthcare
        self.assertEqual(v["median"], 22000.0)                              # commercial: exchange 20401.30, UHC 22000, Anthem 26000 (Medicare left out)
        self.assertEqual(v["medicare"], T.ipps_total(2.0))
        self.assertEqual(v["cost"], round(58612 * T.CCR, 2))
        self.assertEqual(r.totals["payer_rate"], 22000.0)
        self.assertIn("drg_check", {f["kind"] for f in r.flags})

    def test_service_date_before_the_price_file_gives_an_estimate(self):
        b = reader.read(bill("B01"), "b.pdf")
        b.date_from = b.date_to = date(2026, 3, 1)
        for ln in b.lines:
            ln.service_date = date(2026, 3, 1)
        con = db.connect()
        r = price.price(con, b, "360085", b.payer)
        con.close()
        self.assertFalse(r.price_file.covers_service_date)
        self.assertIsNone(r.lines["service_date"][0]["payer_rate"])
        self.assertEqual(r.lines["estimate"][0]["payer_rate"], Decimal("1450.00"))
        self.assertTrue(r.lines["estimate"][0]["payer_rate_is_estimate"])
        self.assertIn("estimate", {f["kind"] for f in r.flags})
        self.assertEqual(r.totals["basis"], "estimate")                    # the report uses the nearest file ...
        self.assertEqual(r.totals["payer_rate"], 1450.0 + 825.6 + 200.0 + 20.0)   # no OPPS rate in March, so no %-of-Medicare value
        self.assertTrue(r.totals["payer_rate_is_estimate"])                 # ... and says every value is an estimate
        self.assertNotIn("Line 1", " ".join(f["title"] for f in r.flags if f["kind"] == "no_rate"))

    def test_unknown_insurer_gets_no_rate_and_a_reversal_line_subtracts(self):
        b = reader.read(bill("B01"), "b.pdf")
        con = db.connect()
        r = price.price(con, b, "360085", "XYZ Benefit Trust")
        self.assertEqual(r.payer_group, "Other")
        self.assertTrue(all(l["payer_rate"] is None for l in r.lines["service_date"]))   # never another plan's rate
        self.assertIn("not in the payer list", " ".join(f["title"] for f in r.flags))
        import copy
        b2 = copy.deepcopy(b)
        rev = copy.deepcopy(b2.lines[0]); rev.line_no = 12; rev.charge = -rev.charge; b2.lines.append(rev)
        r2 = price.price(con, b2, "360085", "Anthem")
        con.close()
        self.assertEqual(r2.lines["service_date"][11]["payer_rate"], Decimal("-1450.00"))
        self.assertEqual(r2.totals["payer_rate"], r.totals["billed"] and round(825.6 + 176.76 + 20.0, 2))
        self.assertEqual(r2.totals["billed"], round(9647.87 - 3403, 2))

    def test_price_file_without_payer_group_column(self):
        # Files loaded with the first load_mrf.py have no payer_group: it is worked out from payer and plan names.
        import shutil, duckdb
        d = Path(tempfile.mkdtemp())
        shutil.copy(DATA / "meshkor.duckdb", d / "meshkor.duckdb")
        c = duckdb.connect(str(d / "meshkor.duckdb")); c.execute("ALTER TABLE ref_hospital_mrf DROP COLUMN payer_group"); c.close()
        b = reader.read(bill("B01"), "b.pdf")
        con = db.connect(d / "meshkor.duckdb")
        r = price.price(con, b, "360085", b.payer)
        con.close()
        self.assertEqual(r.lines["service_date"][0]["payer_rate"], Decimal("1450.00"))

    def test_patient_on_a_government_plan_still_gets_its_own_rate(self):
        # The median is commercial only, but a Medicare Advantage patient's own plan rate is still shown.
        b = reader.read(bill("B01"), "b.pdf")
        con = db.connect()
        r = price.price(con, b, "360085", "Humana Medicare Advantage")
        con.close()
        er = r.lines["service_date"][0]
        self.assertEqual(r.payer_group, "Medicare Advantage")
        self.assertEqual(er["payer_rate"], Decimal("620.00"))
        self.assertEqual(er["median"], Decimal("1450.00"))

    def test_hospital_match_by_npi_and_by_name(self):
        con = db.connect()
        self.assertEqual(match.find(con, npi="1124127683").ccn, "360242")
        m = match.find(con, name="Ohio State University Hospital")
        self.assertEqual((m.ccn, m.method, m.certain), ("360085", "name", False))
        self.assertIsNone(match.find(con, name="Nowhere General"))
        con.close()


@override_settings(MESHKOR_FILE_KEY=Fernet.generate_key().decode(), MESHKOR_UPLOAD_DIR=TMP)
@mock.patch.dict(os.environ, {"MESHKOR_DATA_ROOT": str(DATA)})
class ApiBase(APITestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.firm = Firm.objects.create(name="Smith & Lee LLP")
        cls.rev = User.objects.create_user(username="rev@mk.test", email="rev@mk.test", password="pw-12345")
        Profile.objects.create(user=cls.rev, role="reviewer")

    def setUp(self):
        r = self.client.post("/api/auth/login/", {"email": "rev@mk.test", "password": "pw-12345"}, format="json")
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {r.data['access']}")

    def upload(self, name, eob_name=None, eob_bytes=None, eob_file="eob.pdf"):
        d = {"report_for": "plaintiff", "consent": "true", "firm_id": self.firm.id,
             "bill": SimpleUploadedFile(f"{name}.pdf", bill(name), content_type="application/pdf")}
        if eob_name or eob_bytes:
            d["eob"] = SimpleUploadedFile(eob_file, eob_bytes or eob(eob_name))
        return self.client.post("/api/cases/", d, format="multipart")


@override_settings(MESHKOR_FILE_KEY=Fernet.generate_key().decode(), MESHKOR_UPLOAD_DIR=TMP)
@mock.patch.dict(os.environ, {"MESHKOR_DATA_ROOT": str(DATA)})
class IntakeApiTests(ApiBase):

    def test_upload_prices_the_bill(self):
        r = self.upload("B01")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.data["status"], "needs_review")
        self.assertEqual(r.data["hospital_name"], "Ohio State University Hospitals")
        self.assertEqual(r.data["billed_total"], 9647.87)
        cid = r.data["id"]
        det = self.client.get(f"/api/cases/{cid}/").data
        self.assertEqual(det["hospital_match"]["ccn"], "360085")
        self.assertTrue(det["hospital_match"]["price_file"]["covers_service_date"])
        self.assertEqual(det["totals"]["medicare"], Case.objects.get(pk=cid).totals["medicare"])
        self.assertIn("no_eob", {f["kind"] for f in det["flags"]})
        lines = self.client.get(f"/api/cases/{cid}/lines/").data
        self.assertEqual(len(lines), 11); self.assertEqual(lines[0]["payer_rate"], 1450.0)
        self.assertEqual(lines[0]["sources"][0]["label"], "Medicare cost report, year ending 2024-06-30")
        rep = self.client.get(f"/api/cases/{cid}/report/").data
        pr = rep["summary"][1]
        self.assertEqual(pr["key"], "payer_rate")
        cov = Case.objects.get(pk=cid).totals["coverage"]
        self.assertLess(cov["payer_rate"], 100)                           # the Anthem total leaves out some lines ...
        self.assertEqual(pr["pct_of_billed"], round(100 * pr["total"] / cov["payer_rate_billed"]))   # ... so its % uses those lines only
        self.assertIn("lines", pr["note"])
        self.assertIn(f"{pr['pct_of_billed']} % of the billed charges on those", rep["summary_text"])  # text and table agree
        self.assertTrue(all(r["key"] != "cash" or r["pct_of_billed"] is None for r in rep["summary"]))  # cash on 1 line: not compared
        pdf = self.client.get(f"/api/cases/{cid}/report.pdf")
        self.assertEqual(pdf.status_code, 200); self.assertEqual(pdf.content[:4], b"%PDF")

    def test_inpatient_upload_and_drg_change_reprices(self):
        cid = self.upload("B08").data["id"]
        det = self.client.get(f"/api/cases/{cid}/").data
        rep = self.client.get(f"/api/cases/{cid}/report/").data
        self.assertEqual(rep["stay"]["ms_drg"], "481"); self.assertIn("MS-DRG 481", rep["summary_text"])
        self.assertTrue(any(s["data"].startswith("Medicare IPPS") for s in rep["sources"]))
        self.assertEqual(rep["summary"][1]["pct_of_billed"], round(100 * 22000 / 58612))
        self.assertEqual(det["setting"], "inpatient"); self.assertEqual(det["stay"]["ms_drg"], "481")
        self.assertEqual(det["stay"]["values"]["medicare"], T.ipps_total(2.0))
        r = self.client.post(f"/api/cases/{cid}/drg/confirm/", {"ms_drg": "480"}, format="json")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual((r.data["ms_drg"], r.data["drg_source"], r.data["drg_confirmed"]), ("480", "reviewer", True))
        self.assertEqual(r.data["values"]["medicare"], T.ipps_total(3.0))
        self.assertFalse(Case.objects.get(pk=cid).flags.filter(kind="drg_check", resolved=False).exists())

    def test_hospital_without_price_file(self):
        r = self.upload("B10")
        self.assertEqual(r.data["status"], "no_price_file")
        rep = self.client.get(f"/api/cases/{r.data['id']}/report/").data
        self.assertEqual(rep["summary"][1]["benchmark"], "No insurer on the bill")   # self-pay, no made-up insurer name
        lines = self.client.get(f"/api/cases/{r.data['id']}/lines/").data
        self.assertEqual(lines[0]["medicare"], 110.0)                       # national OPPS: James has no IPPS wage index

    def test_photo_upload_is_priced_and_flagged_for_checking(self):
        photo = ReaderAndPriceTests._photo(self, "B01")
        r = self.client.post("/api/cases/", {"report_for": "plaintiff", "consent": "true", "firm_id": self.firm.id,
                                             "bill": SimpleUploadedFile("bill.jpg", photo, content_type="image/jpeg")}, format="multipart")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.data["status"], "needs_review")
        det = self.client.get(f"/api/cases/{r.data['id']}/").data
        self.assertEqual(det["bill_type"], "UB-04 (OCR)")
        f = det["flags"][0]
        self.assertEqual((f["kind"], f["severity"]), ("low_confidence", "check"))
        self.assertIn("by OCR", f["title"])

    def test_renamed_file_is_refused(self):
        r = self.client.post("/api/cases/", {"report_for": "plaintiff", "consent": "true", "firm_id": self.firm.id,
                                             "bill": SimpleUploadedFile("bill.pdf", b"MZ\x90\x00 not a pdf", content_type="application/pdf")},
                             format="multipart")
        self.assertEqual(r.status_code, 400); self.assertIn("not a real PDF", str(r.data))

    def test_unreadable_bill_goes_to_error(self):
        r = self.client.post("/api/cases/", {"report_for": "plaintiff", "consent": "true", "firm_id": self.firm.id,
                                             "bill": SimpleUploadedFile("scan.pdf", b"%PDF-1.4\n%%EOF", content_type="application/pdf")},
                             format="multipart")
        self.assertEqual(r.status_code, 201)
        self.assertEqual(r.data["status"], "error"); self.assertIn("could not be read", r.data["status_detail"])

    # ---- background pricing (cases/processing.py) ----
    def test_worker_mode_queues_then_the_worker_prices(self):
        from io import StringIO
        from django.core.management import call_command
        with override_settings(MESHKOR_PRICING="worker"):
            r = self.upload("B01")
            self.assertEqual(r.status_code, 201)
            self.assertEqual(r.data["status"], "uploaded")                 # the upload answers at once
            det = self.client.get(f"/api/cases/{r.data['id']}/").data
            self.assertEqual(det["status"], "uploaded")
            out = StringIO()
            call_command("pricing_worker", "--once", stdout=out)
        self.assertIn("needs_review", out.getvalue())
        self.assertIn("1 case(s) priced", out.getvalue())
        c = Case.objects.get(pk=r.data["id"])
        self.assertEqual((c.status, float(c.billed_total)), ("needs_review", 9647.87))
        out = StringIO(); call_command("pricing_worker", "--once", stdout=out)
        self.assertIn("0 case(s) priced", out.getvalue())                 # never priced twice

    def test_a_case_is_claimed_by_one_caller_only(self):
        from . import processing
        with override_settings(MESHKOR_PRICING="worker"):
            cid = self.upload("B01").data["id"]
        self.assertTrue(processing.claim(cid))
        self.assertFalse(processing.claim(cid))
        self.assertFalse(processing.price_waiting_case(cid))

    def test_a_case_stuck_by_a_server_stop_is_queued_again(self):
        from datetime import timedelta
        from django.utils import timezone
        from . import processing
        with override_settings(MESHKOR_PRICING="worker"):
            cid = self.upload("B01").data["id"]
        processing.claim(cid)                                                # it started, then the server stopped
        Case.objects.filter(pk=cid).update(updated_at=timezone.now() - timedelta(minutes=31))
        self.assertIn(cid, [c.pk for c in processing.waiting()])
        self.assertEqual(Case.objects.get(pk=cid).status, "uploaded")
        self.assertTrue(processing.price_waiting_case(cid))
        self.assertEqual(Case.objects.get(pk=cid).status, "needs_review")

    def test_thread_mode_prices_after_the_upload_is_saved(self):
        from . import processing
        with override_settings(MESHKOR_PRICING="thread"), \
             mock.patch.object(processing, "_submit", side_effect=lambda cid, uid: processing.price_waiting_case(cid, uid)) as sub, \
             self.captureOnCommitCallbacks(execute=True):
            r = self.upload("B01")
        self.assertEqual(r.data["status"], "uploaded")
        sub.assert_called_once()
        self.assertEqual(Case.objects.get(pk=r.data["id"]).status, "needs_review")

    def test_thread_mode_resumes_cases_left_by_a_restart(self):
        from datetime import timedelta
        from django.utils import timezone
        from . import processing
        with override_settings(MESHKOR_PRICING="worker"):
            cid = self.upload("B01").data["id"]
        Case.objects.filter(pk=cid).update(updated_at=timezone.now() - timedelta(minutes=3))
        with override_settings(MESHKOR_PRICING="thread"), \
             mock.patch.object(processing, "_submit", side_effect=lambda c, u: processing.price_waiting_case(c, u)):
            self.client.get("/api/cases/")                                  # a reviewer opens the case list
        self.assertEqual(Case.objects.get(pk=cid).status, "needs_review")


class EobReaderTests(SimpleTestCase):
    """pipeline/lien/eob.py on the 3 synthetic EOB layouts."""

    def test_anthem_layout_line_by_line(self):
        from pipeline.lien import eob as E
        e = E.read(eob("B01"))
        self.assertEqual((e.claim_number, e.provider, str(e.date_from)), ("TEST-CLM-B01", "OHIO STATE UNIVERSITY HOSPITALS", "2026-08-14"))
        self.assertIn("Anthem", e.payer)
        self.assertEqual(e.columns, ["billed", "allowed", "deductible", "copay", "coinsurance", "paid", "patient"])
        self.assertEqual((str(e.total("allowed")), str(e.total("paid")), str(e.total("patient"))), ("3501.80", "2401.44", "1100.36"))
        self.assertEqual(e.total("allowed"), e.stated_totals["allowed"])          # the lines add up to the EOB's own total
        b = reader.read(bill("B01"), "b.pdf")
        m = E.match(e, b.lines)
        self.assertEqual(len(m), 11)
        self.assertEqual(str(m[1].get("allowed")), "1450.00")

    def test_other_column_order_and_two_digit_years(self):
        from pipeline.lien import eob as E
        e = E.read(eob("B07"))
        self.assertEqual(e.columns, ["billed", "discount", "allowed", "not_covered", "deductible", "coinsurance", "paid", "patient"])
        self.assertEqual(str(e.date_from), "2026-08-20")
        m = E.match(e, reader.read(bill("B07"), "b.pdf").lines)
        self.assertEqual(len(m), 7)                                               # lines with only a revenue code match too
        self.assertEqual(str(m[4].get("allowed")), "2900.00")                     # 0278 implant
        self.assertEqual(m[2].revenue_code, "0370")

    def test_inpatient_eob_is_compared_as_one_stay(self):
        from pipeline.lien import eob as E
        e = E.read(eob("B08"))
        self.assertEqual((e.ms_drg, str(e.total("allowed"))), ("481", "47909.59"))
        s, per_line, flags = E.compare(e, reader.read(bill("B08"), "b.pdf"), "UnitedHealthcare", payer_rate_total=47909.59)
        self.assertEqual((s["level"], per_line), ("stay", {}))
        self.assertEqual(s["vs_price_file"]["pct"], 0)
        self.assertFalse([f for f in flags if f["severity"] == "check"])

    def test_a_bill_is_not_an_eob(self):
        from pipeline.lien import eob as E
        with self.assertRaises(E.EobNotReadable):
            E.read(bill("B01"))

    def test_eob_for_another_claim_is_flagged(self):
        from pipeline.lien import eob as E
        s, _pl, flags = E.compare(E.read(eob("B08")), reader.read(bill("B01"), "b.pdf"), "Anthem")
        titles = " | ".join(f["title"] for f in flags if f["severity"] == "check")
        self.assertIn("EOB is from UnitedHealthcare, the bill names Anthem", titles)
        self.assertIn("dates of service do not overlap", titles)
        self.assertIn("billed total differs", titles)

    def test_photo_of_an_eob_is_read_by_ocr(self):
        import pypdfium2 as pdfium
        from io import BytesIO
        from pipeline.lien import eob as E
        img = pdfium.PdfDocument(eob("B01"))[0].render(scale=200 / 72).to_pil()
        buf = BytesIO(); img.save(buf, "PNG")
        e = E.read(buf.getvalue(), "eob.png")
        self.assertEqual(e.method, "ocr")
        self.assertEqual(str(e.total("allowed")), "3501.80")
        self.assertLessEqual(e.ocr_score, 0.95)


@override_settings(MESHKOR_FILE_KEY=Fernet.generate_key().decode(), MESHKOR_UPLOAD_DIR=TMP)
@mock.patch.dict(os.environ, {"MESHKOR_DATA_ROOT": str(DATA)})
class EobApiTests(ApiBase):
    """Bill + EOB through the API, report, CSV and PDF."""

    def test_upload_with_eob_fills_the_report(self):
        r = self.upload("B01", "B01")
        self.assertEqual(r.data["status"], "needs_review")
        cid = r.data["id"]
        det = self.client.get(f"/api/cases/{cid}/").data
        self.assertEqual(det["eob"]["totals"]["allowed"], 3501.8)
        self.assertEqual((det["eob"]["lines_matched"], det["eob"]["bill_lines"]), (11, 11))
        kinds = {f["kind"] for f in det["flags"]}
        self.assertNotIn("no_eob", kinds); self.assertIn("eob", kinds)
        self.assertIn("vs_price_file", det["eob"])
        lines = self.client.get(f"/api/cases/{cid}/lines/").data
        self.assertEqual((lines[0]["eob_allowed"], lines[0]["eob_paid"], lines[0]["eob_patient"]), (1450.0, 760.0, 690.0))
        rep = self.client.get(f"/api/cases/{cid}/report/").data
        keys = [x["key"] for x in rep["summary"]]
        self.assertEqual(keys[:4], ["billed", "eob_allowed", "eob_paid", "payer_rate"])
        ea = rep["summary"][1]
        self.assertEqual((ea["total"], ea["pct_of_billed"]), (3501.8, round(100 * 3501.8 / 9647.87)))
        self.assertIn("allowed $3,501.80 (36 % of billed) and paid $2,401.44; the patient's share was $1,100.36", rep["summary_text"])
        self.assertEqual(rep["eob"]["claim_number"], "TEST-CLM-B01")
        self.assertNotIn("file_name", rep["eob"])
        self.assertIn("Insurer EOB uploaded with the case", [s["data"] for s in rep["sources"]])
        csv_ = self.client.get(f"/api/cases/{cid}/report.csv").content.decode("utf-8-sig")
        self.assertIn("eob_allowed", csv_); self.assertIn("3501.8", csv_)
        pdf = self.client.get(f"/api/cases/{cid}/report.pdf")
        self.assertEqual(pdf.content[:4], b"%PDF")

    def test_inpatient_eob(self):
        cid = self.upload("B08", "B08").data["id"]
        det = self.client.get(f"/api/cases/{cid}/").data
        self.assertEqual((det["eob"]["level"], det["eob"]["ms_drg"]), ("stay", "481"))
        rep = self.client.get(f"/api/cases/{cid}/report/").data
        self.assertEqual(rep["summary"][1]["total"], 47909.59)

    def test_unreadable_eob_does_not_stop_the_case(self):
        r = self.upload("B01", eob_bytes=bill("B07"))                     # a bill put in the EOB box
        self.assertEqual(r.data["status"], "needs_review")
        det = self.client.get(f"/api/cases/{r.data['id']}/").data
        self.assertIsNone(det["eob"])
        f = [x for x in det["flags"] if x["title"] == "The EOB could not be read"]
        self.assertEqual(f[0]["severity"], "check")
        rep = self.client.get(f"/api/cases/{r.data['id']}/report/").data
        self.assertNotIn("eob_allowed", [x["key"] for x in rep["summary"]])


@override_settings(MESHKOR_FILE_KEY=Fernet.generate_key().decode(), MESHKOR_UPLOAD_DIR=TMP)
@mock.patch.dict(os.environ, {"MESHKOR_DATA_ROOT": str(DATA)})
class ReportPdfTests(ApiBase):
    """The designed PDF (cases/report_pdf.py). Set MESHKOR_SAVE_PDF=<folder> to keep the files and look at them."""

    def pdf(self, cid):
        r = self.client.get(f"/api/cases/{cid}/report.pdf")
        self.assertEqual((r.status_code, r["Content-Type"]), (200, "application/pdf"))
        out = os.environ.get("MESHKOR_SAVE_PDF")
        return r.content, out

    @staticmethod
    def text(data):
        import io
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as p:
            return [pg.extract_text() or "" for pg in p.pages]

    def approve(self, cid):
        from .models import Flag
        Flag.objects.filter(case_id=cid).update(resolved=True)
        self.assertEqual(self.client.post(f"/api/cases/{cid}/approve/", {"note": ""}, format="json").status_code, 200)

    def test_draft_and_final_outpatient_with_eob(self):
        cid = self.upload("B01", "B01").data["id"]
        data, out = self.pdf(cid)
        pages = self.text(data)
        first = pages[0]
        for s in ("MeshKor Lien Valuation", "Reasonable value of hospital charges", "MK-2026-", "Smith & Lee LLP",
                  "Ohio State University Hospitals", "1. Summary", "$3,501.80", "Allowed by Anthem"):
            self.assertIn(s, first, s)
        self.assertIn("DRAFT", first)
        self.assertIn(f"Page 1 of {len(pages)}", first)
        whole = "\n".join(pages)
        for s in ("2. Line by line", "3. Sources", "4. Method", "5. Limits", "Insurer's Explanation of Benefits",
                  "Total (11 lines)", "9,647.87 3,501.80", "not been approved"):
            self.assertIn(s, whole, s)
        if out:
            Path(out, "B01_draft.pdf").write_bytes(data)
        self.approve(cid)
        data, out = self.pdf(cid)
        whole = "\n".join(self.text(data))
        self.assertIn("Reviewed and approved", whole)
        self.assertNotIn("not been approved", whole)
        if out:
            Path(out, "B01_final.pdf").write_bytes(data)

    def test_inpatient_pdf(self):
        cid = self.upload("B08", "B08").data["id"]
        data, out = self.pdf(cid)
        whole = "\n".join(self.text(data))
        self.assertIn("Inpatient stay", whole)
        self.assertIn("MS-DRG: 481", whole)
        if out:
            Path(out, "B08_draft.pdf").write_bytes(data)
