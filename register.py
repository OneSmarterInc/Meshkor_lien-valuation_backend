"""Data register: one Excel file that shows every download, what it feeds,
whether it is loaded into Parquet, and what is still to download.

    python register.py                       # -> data_register.xlsx (in the repo root)
    python register.py --raw "E:\\Datasetof CMS" --out "E:\\data_register.xlsx"

Read-only: it scans the raw folder, the load manifests and data/meshkor.duckdb.
Run it after every download and after every load. The sheets "Download log"
and "Notes" are yours: rows you type there are kept when the file is rebuilt.
"""
from __future__ import annotations

import argparse
import datetime as dt
import glob
import json
import os
import shutil
from collections import defaultdict
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
KEEP_SHEETS = ("Download log", "Notes")
GREEN, AMBER, RED, GREY = "C6EFCE", "FFEB9C", "FFC7CE", "EDEDED"

STEP1 = {"ref_provider", "ref_mpfs", "ref_mpfs_locality", "ref_opps", "ref_ipps_drg", "ref_ipps_hospital",
         "ref_asc", "ref_ncci_ptp", "ref_mue", "ref_hcris", "ref_ccn_npi", "ref_code_specialty", "ref_clfs",
         "ref_dmepos", "ref_dmepos_rural_zip", "ref_ambulance", "ref_zip_locality", "ref_anesthesia_cf",
         "ref_anesthesia_base_units", "ref_asp", "ref_ndc_hcpcs", "ref_hcpcs_code"}


def env_value(key: str) -> str | None:
    if os.environ.get(key):
        return os.environ[key]
    env = ROOT / ".env"
    if env.exists():
        for line in env.read_text(encoding="utf-8").splitlines():
            if line.strip().startswith(key + "="):
                return line.split("=", 1)[1].strip().strip('"')
    return None


def group_of(table: str) -> str:
    return "Step 1" if table in STEP1 else "§11"


def gb(n: float) -> float:
    return round(n / 1e9, 3)


def scan_raw(raw: Path) -> list[dict]:
    out = []
    for p in raw.rglob("*"):
        if p.is_file() and not p.name.endswith(".sha256"):
            st = p.stat()
            rel = p.relative_to(raw).as_posix()
            out.append({"rel": rel, "top": rel.split("/")[0], "name": p.name, "bytes": st.st_size,
                        "modified": dt.datetime.fromtimestamp(st.st_mtime).replace(microsecond=0)})
    return sorted(out, key=lambda r: r["rel"].lower())


def loaded_rows(db: Path) -> tuple[dict, dict]:
    """(table, period) -> rows, and table -> summary, from meshkor.duckdb."""
    per, tab = {}, {}
    if not db.exists():
        return per, tab
    import duckdb
    con = duckdb.connect(str(db), read_only=True)
    for (t,) in con.execute("SELECT table_name FROM duckdb_tables() WHERE table_name LIKE 'ref_%'").fetchall():
        cols = {r[0] for r in con.execute(f"DESCRIBE {t}").fetchall()}
        n = con.execute(f"SELECT count(*) FROM {t}").fetchone()[0]
        rng = con.execute(f"SELECT min(effective_start), max(effective_end) FROM {t}").fetchone() \
            if "effective_start" in cols else (None, None)
        tab[t] = {"rows": n, "from": rng[0], "to": rng[1], "releases": 0}
        if "period" in cols:
            for p, k in con.execute(f"SELECT period, count(*) FROM {t} GROUP BY 1").fetchall():
                per[(t, str(p).split("/")[0])] = per.get((t, str(p).split("/")[0]), 0) + k
            tab[t]["releases"] = con.execute(f"SELECT count(DISTINCT period) FROM {t}").fetchone()[0]
    con.close()
    return per, tab


def manifests(ref: Path) -> dict:
    out = {}
    for f in (ref / "_manifest").glob("*.json"):
        try:
            out[f.stem] = json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            pass
    return out


def size_of(path: Path) -> int:
    if path.is_file():
        return path.stat().st_size
    return sum(p.stat().st_size for p in path.rglob("*") if p.is_file()) if path.exists() else 0


def write_sheet(wb, title, header, rows, widths=None, status_col=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, 1):
        w = (widths or {}).get(h) or min(60, max(10, len(str(h)) + 2,
                                               *[len(str(r[i - 1])) + 2 for r in rows[:300] if r[i - 1] is not None]))
        ws.column_dimensions[get_column_letter(i)].width = w
    if status_col and rows:
        col = get_column_letter(header.index(status_col) + 1)
        rng = f"{col}2:{col}{len(rows) + 1}"
        for word, color in (("Loaded", GREEN), ("Done", GREEN), ("OK", GREEN), ("Downloaded", AMBER),
                            ("Partial", AMBER), ("Not used", GREY), ("Missing", RED), ("To do", RED)):
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal",
                                                          formula=[f'"{word}"'], fill=PatternFill("solid", fgColor=color)))
    return ws


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Build data_register.xlsx")
    ap.add_argument("--raw", default=env_value("MESHKOR_RAW_DIR") or str(ROOT / "data" / "raw"))
    ap.add_argument("--sources", default=env_value("MESHKOR_SOURCES") or "sources_local.yaml")
    ap.add_argument("--data", default=str(ROOT / "data"))
    ap.add_argument("--out", default=str(ROOT / "data_register.xlsx"))
    a = ap.parse_args(argv)
    raw, data, out = Path(a.raw), Path(a.data), Path(a.out)
    cfg = yaml.safe_load(open(ROOT / "config" / a.sources, encoding="utf-8"))
    planned = yaml.safe_load(open(ROOT / "config" / "register_planned.yaml", encoding="utf-8")) or []
    files = scan_raw(raw) if raw.exists() else []
    per_rows, tab = loaded_rows(data / "meshkor.duckdb")
    mans = manifests(data / "ref")

    # releases -> folders, and file -> releases
    rel_rows, used_by = [], defaultdict(set)
    for t, spec in cfg.items():
        for r in spec.get("releases", []):
            per = str(r["period"])
            label = per + (f"/{r['setting']}" if r.get("setting") else "")
            f0 = r.get("file")
            choices = f0 if isinstance(f0, list) else None          # `file:` may list choices; first found wins
            paths = r.get("files") or ([next((c for c in choices if (raw / c).exists()), choices[0])] if choices else [f0])
            fl = [f for f in files if any(f["rel"] == p or f["rel"].startswith(p.rstrip("/") + "/") for p in paths if p)]
            for f in fl:
                used_by[f["rel"]].add(f"{t} {label}")
            present = all((raw / p).exists() for p in paths if p)
            n = per_rows.get((t, per))
            if n is None and (t, per) not in per_rows and t in tab and not tab[t]["releases"]:
                n = tab[t]["rows"]
            src = next((s for s in mans.get(t, {}).get("sources", []) if s.get("period") == label), {})
            status = "Missing" if not present else ("Loaded" if n else "Downloaded")
            rel_rows.append([group_of(t), t, label, str(r.get("effective_start", "")), str(r.get("effective_end", "")),
                             ", ".join(p for p in paths if p), "Y" if present else "N", len(fl),
                             gb(sum(f["bytes"] for f in fl)), n, (src.get("sha256") or "")[:16],
                             str(mans.get(t, {}).get("finished_at", ""))[:10], r.get("page", ""), status])

    # files sheet
    file_rows = [[f["top"], f["rel"], round(f["bytes"] / 1e6, 2), f["modified"],
                  "; ".join(sorted(used_by[f["rel"]])) or "—",
                  "Loaded" if any(per_rows.get((u.split(" ")[0], u.split(" ")[1].split("/")[0])) for u in used_by[f["rel"]])
                  else ("Downloaded" if used_by[f["rel"]] else "Not used")] for f in files]

    # planned sheet
    plan_rows = []
    for p in planned:
        found = len(glob.glob(str(raw / p["pattern"]), recursive=True)) if p.get("pattern") else 0
        exp = int(p.get("expected") or 0)
        st = "To do" if not exp else ("Done" if found >= exp else ("Partial" if found else "Missing"))
        plan_rows.append([p.get("priority"), p.get("group"), p.get("item"), p.get("why"), p.get("link"),
                          p.get("pattern"), exp, found, p.get("owner"), st])
    plan_rows.sort(key=lambda r: (r[0] or 9, r[1]))

    # parquet sheet
    pq_rows = []
    for t in sorted(set(tab) | set(mans)):
        m = mans.get(t, {})
        p = data / m.get("output", f"ref/{t}.parquet")
        pq_rows.append([group_of(t), t, str(p.relative_to(data)) if p.exists() else "missing",
                        round(size_of(p) / 1e6, 1), tab.get(t, {}).get("rows"), tab.get(t, {}).get("releases"),
                        str(tab.get(t, {}).get("from") or ""), str(tab.get(t, {}).get("to") or ""),
                        str(m.get("finished_at", ""))[:19], "Loaded" if t in tab else "Downloaded"])

    # summary
    st_count = defaultdict(lambda: defaultdict(int))
    for r in rel_rows:
        st_count[r[0]][r[-1]] += 1
    top_size = defaultdict(int)
    for f in files:
        top_size[f["top"]] += f["bytes"]
    summary = [["Generated", dt.datetime.now().replace(microsecond=0)], ["Raw folder", str(raw)],
               ["Data folder (Parquet)", str(data / "ref")], [], ["CONFIGURED RELEASES", "Loaded", "Downloaded, not loaded", "Missing"]]
    for g in sorted(st_count):
        summary.append([g, st_count[g]["Loaded"], st_count[g]["Downloaded"], st_count[g]["Missing"]])
    summary += [[], ["PLANNED (not configured yet)", "Done", "Partial", "Missing / to do"]]
    pc = defaultdict(int)
    for r in plan_rows:
        pc[r[-1]] += 1
    summary.append(["All planned items", pc["Done"], pc["Partial"], pc["Missing"] + pc["To do"]])
    summary += [[], ["SIZE", "GB"], ["Raw downloads (all)", gb(sum(top_size.values()))],
                ["Parquet (data/ref)", gb(size_of(data / "ref"))], ["DuckDB file (rebuilt from Parquet)", gb(size_of(data / "meshkor.duckdb"))],
                ["Raw files not used by any loader", gb(sum(f["bytes"] for f in files if not used_by[f["rel"]]))]]
    summary += [[], ["DISK", "Free GB", "Total GB"]]
    for label, path in (("Raw folder drive", raw), ("Data folder drive", data)):
        try:
            u = shutil.disk_usage(path.anchor or path)
            summary.append([f"{label} ({path.anchor})", gb(u.free), gb(u.total)])
        except Exception:
            pass
    summary += [[], ["RAW SIZE BY FOLDER", "GB", "Files"]]
    cnt = defaultdict(int)
    for f in files:
        cnt[f["top"]] += 1
    for k in sorted(top_size, key=lambda k: -top_size[k]):
        summary.append([k, gb(top_size[k]), cnt[k]])

    # keep the user's sheets
    kept = {}
    if out.exists():
        try:
            old = load_workbook(out)
            for name in KEEP_SHEETS:
                if name in old.sheetnames:
                    kept[name] = [list(r) for r in old[name].iter_rows(values_only=True)]
        except Exception:
            pass

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for r in summary:
        ws.append(r)
    ws.column_dimensions["A"].width = 40
    for c in "BCD":
        ws.column_dimensions[c].width = 22
    for row in ws.iter_rows():
        if row[0].value and str(row[0].value).isupper():
            for c in row:
                c.font = Font(bold=True)
    write_sheet(wb, "Releases", ["Group", "Table", "Release", "From", "To", "Raw folder / file", "Raw present",
                                 "Files", "Raw GB", "Loaded rows", "SHA-256 (first 16)", "Loaded on", "CMS page", "Status"],
                rel_rows, {"CMS page": 50, "Raw folder / file": 45}, "Status")
    write_sheet(wb, "Planned", ["Priority", "Group", "Item", "Why", "Link", "Search pattern", "Expected files",
                                "Found files", "Owner", "Status"], plan_rows, {"Item": 50, "Why": 45, "Link": 50}, "Status")
    write_sheet(wb, "Parquet", ["Group", "Table", "Parquet path", "Size MB", "Rows", "Releases", "From", "To",
                                "Last load", "Status"], pq_rows, None, "Status")
    write_sheet(wb, "Files", ["Folder", "Relative path", "Size MB", "Modified", "Used by (table release)", "Status"],
                file_rows, {"Relative path": 80, "Used by (table release)": 45}, "Status")
    for name in KEEP_SHEETS:
        rows = kept.get(name) or [["Date", "Who", "File / item", "Source page or link", "Raw folder", "Note"]]
        s = wb.create_sheet(name)
        for r in rows:
            s.append(r)
        for c in s[1]:
            c.font = Font(bold=True)
        for i in range(1, 7):
            s.column_dimensions[get_column_letter(i)].width = 30
    wb.save(out)
    print(f"written: {out}")
    for r in summary[4:4 + len(st_count) + 1]:
        print("  ", *r)
    print(f"   planned: {pc['Done']} done, {pc['Partial']} partial, {pc['Missing'] + pc['To do']} missing/to do")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
